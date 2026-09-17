# -*- coding: utf-8 -*-
"""
Iteration 数据汇总脚本：
分析退火迭代（0, 5, 10, 25, 100 轮）的细胞数据变化。

使用方法：
    python iteration_summarize.py <根目录路径> [--constraints 约束文件夹名 ...]

输出：
- 每个约束文件夹一个汇总 Excel（含 ellipse/edgeAngle/ME_MA + 统计 + delta_detail + C1/C2拟合）
- 一个跨约束 C1/C2 对比文件
"""
import os
import sys
import argparse
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# 导入 summarize.py 中的共享函数
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from summarize import (
    HEADER_FONT, HEADER_FILL, GROUP_FILL, TITLE_FONT, DELTA_FILL, PRED_FILL, MN_FILL,
    CENTER, LEFT, THIN, BORDER,
    LENGTH_LABELS, ANGLE_LABELS,
    length_bin, angle_bin, style_header,
    split_mc_ic, split_three_way,
    write_raw_sheet, write_overall, write_by_n, write_n_dist_separate,
    write_delta_by_n, write_dist_separate,
)

# ---------- 常量 ----------
ITERATIONS = [0, 5, 10, 25, 100]
PHASE_MAP = {0: 'initial', 5: 'rounds', 10: 'rounds', 25: 'rounds', 100: 'rounds'}

# Delta 对: (名称, 终点迭代, 起点迭代)
DELTA_PAIRS = [
    ('5-0', 5, 0),
    ('10-0', 10, 0),
    ('25-0', 25, 0),
    ('100-0', 100, 0),
    ('10-5', 10, 5),
    ('25-10', 25, 10),
    ('100-25', 100, 25),
]


# ---------- 读取 ----------
def read_iter_ellipse(base, constraint, runs):
    """读取所有迭代的 ellipse 数据。返回 {iter: DataFrame}。"""
    out = {}
    for it in ITERATIONS:
        phase = PHASE_MAP[it]
        frames = []
        for run in runs:
            path = os.path.join(base, constraint, f'run{run}', f'{it}_{phase}_ellipse.xlsx')
            if not os.path.exists(path):
                continue
            df = pd.read_excel(path)
            df = df.rename(columns={
                '细胞序号': 'No.', '长半轴': 'a', '短半轴': 'b',
                '相邻细胞边数和': 'mn', '细胞边数': 'n',
                '细胞周长': 'P', '细胞面积': 'A'
            })
            df = df[['No.', 'a', 'b', 'mn', 'n', 'P', 'A', '当前细胞层']].copy()
            df.insert(0, 'run', run)
            df['No.'] = df['No.'].astype(int)
            df['当前细胞层'] = df['当前细胞层'].astype(int)
            frames.append(df)
        if frames:
            out[it] = pd.concat(frames, ignore_index=True)
        else:
            out[it] = pd.DataFrame(columns=['run', 'No.', 'a', 'b', 'mn', 'n', 'P', 'A', '当前细胞层'])
    return out


def read_iter_edgeangle(base, constraint, runs):
    """读取所有迭代的 edgeAngle 数据。"""
    out = {}
    for it in ITERATIONS:
        phase = PHASE_MAP[it]
        frames = []
        for run in runs:
            path = os.path.join(base, constraint, f'run{run}', f'{it}_{phase}_edgeAngle.xlsx')
            if not os.path.exists(path):
                continue
            df = pd.read_excel(path)
            df = df.rename(columns={
                '细胞序号': 'No.', '边数': 'n', '相邻细胞边数和': 'mn'
            })
            df = df[['No.', 'n', 'mn', '内角', '夹边1', '当前细胞层']].copy()
            df.insert(0, 'run', run)
            df['No.'] = df['No.'].astype(int)
            df['当前细胞层'] = df['当前细胞层'].astype(int)
            frames.append(df)
        if frames:
            out[it] = pd.concat(frames, ignore_index=True)
        else:
            out[it] = pd.DataFrame(columns=['run', 'No.', 'n', 'mn', '内角', '夹边1', '当前细胞层'])
    return out


def read_iter_mema(base, constraint, runs):
    """读取所有迭代的 ME_MA 数据。"""
    out = {}
    for it in ITERATIONS:
        phase = PHASE_MAP[it]
        frames = []
        for run in runs:
            path = os.path.join(base, constraint, f'run{run}', f'{it}_{phase}_ME_MA.xlsx')
            if not os.path.exists(path):
                continue
            df = pd.read_excel(path)
            df = df.rename(columns={
                '细胞序号': 'No.', '边缘边长(ME)': 'ME',
                '边缘角1(MA1)': 'MA1', '边缘角2(MA2)': 'MA2'
            })
            df = df[['No.', 'ME', 'MA1', 'MA2', '当前细胞层']].copy()
            df.insert(0, 'run', run)
            df['No.'] = df['No.'].astype(int)
            df['当前细胞层'] = df['当前细胞层'].astype(int)
            frames.append(df)
        if frames:
            out[it] = pd.concat(frames, ignore_index=True)
        else:
            out[it] = pd.DataFrame(columns=['run', 'No.', 'ME', 'MA1', 'MA2', '当前细胞层'])
    return out


# ---------- Delta 计算 ----------
def compute_delta_iter(ell_from, ell_to, from_it, to_it):
    """计算两个迭代点之间的 delta。返回 DataFrame（含当前细胞层）。"""
    init = ell_from[['run', 'No.', 'n', 'P', 'A', '当前细胞层']].copy()
    fin = ell_to[['run', 'No.', 'n', 'P', 'A']].copy()
    init = init.rename(columns={'P': 'P_from', 'A': 'A_from', 'n': 'n_from'})
    fin = fin.rename(columns={'P': 'P_to', 'A': 'A_to', 'n': 'n_to'})
    m = init.merge(fin, on=['run', 'No.'], how='inner')
    m['ΔP'] = m['P_to'] - m['P_from']
    m['ΔA'] = m['A_to'] - m['A_from']
    m['n'] = m['n_from']
    m['SI_from'] = m['P_from'] / np.sqrt(m['A_from'].clip(lower=1e-12))
    m['SI_to'] = m['P_to'] / np.sqrt(m['A_to'].clip(lower=1e-12))
    m['ΔSI'] = m['SI_to'] - m['SI_from']
    m['ΔEavg'] = m['ΔP'] / m['n'].clip(lower=1)
    return m


def fit_c1_c2_iter(ell_from, ell_to, from_it, to_it):
    """
    拟合 C1 和 C2 (MC/IC 分开)。
    MC: ΔA = C1*(n-5) + C2*(1-A/An)
    IC: ΔA = C1*(n-6) + C2*(1-A/An)
    A 为起点迭代的 A，An 为每个 n 的起点 A 平均值。
    """
    init = ell_from[['run', 'No.', 'n', 'A', '当前细胞层']].copy()
    fin = ell_to[['run', 'No.', 'A']].copy()
    init = init.rename(columns={'A': 'A_from'})
    fin = fin.rename(columns={'A': 'A_to'})
    m = init.merge(fin, on=['run', 'No.'], how='inner')
    m['ΔA'] = m['A_to'] - m['A_from']

    results = {}
    for (mask, label, n_offset) in [
        (m['当前细胞层'] == 1, 'MC', 5),
        (m['当前细胞层'] != 1, 'IC', 6)
    ]:
        sub = m[mask].copy()
        if len(sub) == 0:
            results[label] = {'C1': 0, 'C2': 0, 'RMSE': 0, 'n_cells': 0,
                               'An': {}, 'details': pd.DataFrame()}
            continue

        n_arr = sub['n'].values.astype(float)
        da_arr = sub['ΔA'].values.astype(float)
        a_arr = sub['A_from'].values.astype(float)

        # An: 每个 n 的 A_from 平均值
        an_dict = {}
        for nv in sorted(sub['n'].unique()):
            an_dict[int(nv)] = float(sub[sub['n'] == nv]['A_from'].mean())

        # C1 拟合
        dn = n_arr - n_offset
        mask_safe = dn != 0
        dn_safe = dn[mask_safe]
        da_safe = da_arr[mask_safe]
        if len(dn_safe) > 0:
            c1 = float(np.sum(dn_safe * da_safe) / np.sum(dn_safe ** 2))
        else:
            c1 = 0.0

        # C2 遍历 [0, 3]
        best_c2 = 0.0
        best_rmse = float('inf')
        correction = np.array([1 - a_arr[i] / an_dict.get(int(n_arr[i]), 1.0) for i in range(len(n_arr))])
        for c2_val in np.arange(0.0, 3.001, 0.01):
            delta_a_pred = c1 * dn + c2_val * correction
            residuals = da_arr - delta_a_pred
            rmse = float(np.sqrt(np.mean(residuals ** 2)))
            if rmse < best_rmse:
                best_rmse = rmse
                best_c2 = float(c2_val)

        # 详细数据
        details = sub[['run', 'No.', 'n', 'ΔA', 'A_from']].copy()
        details['An'] = details['n'].map(lambda x: an_dict.get(int(x), 0))
        details['(1-A/An)'] = 1 - details['A_from'] / details['An'].clip(lower=1e-12)
        details['C1'] = c1
        details['C2'] = best_c2
        details['ΔA_pred'] = c1 * dn + best_c2 * details['(1-A/An)'].values
        details['残差'] = details['ΔA'] - details['ΔA_pred']

        results[label] = {
            'C1': c1, 'C2': best_c2, 'RMSE': best_rmse,
            'n_cells': len(sub), 'An': an_dict, 'details': details
        }
    return results


# ---------- Delta detail 写入 ----------
def write_delta_detail_sheet(ws, all_deltas, layer_data_0):
    """delta_detail 工作表：MC/IC → delta值 → delta对，3层表头，逐行原始数据。

    布局:
    Row1: [MC (28列)                                    ] [IC (28列)                                    ]
    Row2: [ΔP(7)     |ΔA(7)     |ΔSI(7)    |ΔEavg(7)   ] [ΔP(7)     |ΔA(7)     |ΔSI(7)    |ΔEavg(7)   ]
    Row3: [5-0|10-0|25-0|100-0|10-5|25-10|100-25 |...  ] [5-0|10-0|...                                ]
    Row4+: 数据行（MC/IC 各自按 run+No. 外连接对齐）
    """
    delta_metrics = ['ΔP', 'ΔA', 'ΔSI', 'ΔEavg']
    pair_names = [p[0] for p in all_deltas]
    n_pairs = len(pair_names)
    n_metrics = len(delta_metrics)
    cols_per_group = n_pairs * n_metrics  # 28

    # 按 MC/IC 分组，每组内按 run+No. 外连接所有 delta对
    merged_data = {}
    for (mc_label, mask_fn) in [('MC', lambda df: df['当前细胞层'] == 1),
                                 ('IC', lambda df: df['当前细胞层'] != 1)]:
        base_df = None
        for pair_name, delta_df in all_deltas:
            sub = delta_df[mask_fn(delta_df)][['run', 'No.', 'ΔP', 'ΔA', 'ΔSI', 'ΔEavg']].copy()
            sub = sub.rename(columns={
                'ΔP': f'ΔP_{pair_name}', 'ΔA': f'ΔA_{pair_name}',
                'ΔSI': f'ΔSI_{pair_name}', 'ΔEavg': f'ΔEavg_{pair_name}'
            })
            if base_df is None:
                base_df = sub
            else:
                base_df = base_df.merge(sub, on=['run', 'No.'], how='outer')
        # 按 run, No. 排序
        base_df = base_df.sort_values(['run', 'No.']).reset_index(drop=True)
        merged_data[mc_label] = base_df

    # 写入 3 层表头 + 数据
    col = 1
    for mc_label in ['MC', 'IC']:
        df = merged_data[mc_label]
        start_col = col
        # Row 1: MC/IC 大标题（合并 28 列）
        ws.merge_cells(start_row=1, start_column=start_col,
                       end_row=1, end_column=start_col + cols_per_group - 1)
        c = ws.cell(row=1, column=start_col, value=mc_label)
        c.font = HEADER_FONT
        c.fill = DELTA_FILL
        c.alignment = CENTER
        for cc in range(start_col, start_col + cols_per_group):
            ws.cell(row=1, column=cc).fill = DELTA_FILL
            ws.cell(row=1, column=cc).border = BORDER

        # Row 2 & 3: delta值（合并7列） + delta对
        for mi, metric in enumerate(delta_metrics):
            metric_start = start_col + mi * n_pairs
            ws.merge_cells(start_row=2, start_column=metric_start,
                           end_row=2, end_column=metric_start + n_pairs - 1)
            c2 = ws.cell(row=2, column=metric_start, value=metric)
            c2.font = HEADER_FONT
            c2.fill = DELTA_FILL
            c2.alignment = CENTER
            for cc in range(metric_start, metric_start + n_pairs):
                ws.cell(row=2, column=cc).fill = DELTA_FILL
                ws.cell(row=2, column=cc).border = BORDER
            # Row 3: delta对名称
            for j, pair_name in enumerate(pair_names):
                style_header(ws.cell(row=3, column=metric_start + j, value=pair_name))
                ws.cell(row=3, column=metric_start + j).fill = DELTA_FILL
            col = metric_start + n_pairs

        # 数据行（从 Row 4 开始）
        for r, (_, row_data) in enumerate(df.iterrows()):
            for mi, metric in enumerate(delta_metrics):
                for k, pair_name in enumerate(pair_names):
                    col_name = f'{metric}_{pair_name}'
                    val = row_data.get(col_name, None)
                    cell = ws.cell(row=4 + r,
                                   column=start_col + mi * n_pairs + k,
                                   value=float(val) if pd.notna(val) else None)
                    cell.border = BORDER
                    if pd.notna(val):
                        cell.number_format = '0.000000'

    # 列宽
    total_cols = 2 * cols_per_group
    for cc in range(1, total_cols + 1):
        ws.column_dimensions[get_column_letter(cc)].width = 11
    ws.freeze_panes = 'A4'


def write_delta_block_c(ws, start_row, title, all_deltas):
    """Block C: Delta 按 n 分组统计，分类同 delta_detail (MC/IC → delta值 → delta对)。

    8 个子表（MC/IC × 4 delta值），每个子表水平展开 7 个 delta对：
    n值 | 5-0均值 | 5-0标准差 | 10-0均值 | 10-0标准差 | ... | 100-25均值 | 100-25标准差
    """
    delta_metrics = ['ΔP', 'ΔA', 'ΔSI', 'ΔEavg']
    pair_names = [p[0] for p in all_deltas]
    n_pairs = len(pair_names)

    ws.cell(row=start_row, column=1, value=title).font = TITLE_FONT
    r = start_row + 1

    for (mc_label, mask_fn) in [('MC', lambda df: df['当前细胞层'] == 1),
                                 ('IC', lambda df: df['当前细胞层'] != 1)]:
        for metric in delta_metrics:
            # 子标题
            ws.cell(row=r, column=1, value=f'{mc_label} - {metric}').font = Font(bold=True)
            r += 1
            # 表头
            headers = ['n值']
            for pair_name in pair_names:
                headers += [f'{pair_name}_均值', f'{pair_name}_标准差']
            for j, h in enumerate(headers):
                style_header(ws.cell(row=r, column=j + 1, value=h))
            r += 1
            # 收集所有 n 值
            all_n = set()
            for pair_name, delta_df in all_deltas:
                sub = delta_df[mask_fn(delta_df)]
                all_n.update(sub['n'].dropna().unique())
            all_n = sorted(all_n)
            # 按 n 分组统计
            for nv in all_n:
                ws.cell(row=r, column=1, value=int(nv)).border = BORDER
                col_idx = 2
                for pair_name, delta_df in all_deltas:
                    sub = delta_df[mask_fn(delta_df) & (delta_df['n'] == nv)][metric].dropna()
                    mean_val = float(sub.mean()) if len(sub) else None
                    std_val = float(sub.std()) if len(sub) else None
                    c1 = ws.cell(row=r, column=col_idx, value=mean_val)
                    c1.border = BORDER
                    if mean_val is not None:
                        c1.number_format = '0.0000'
                    c2 = ws.cell(row=r, column=col_idx + 1, value=std_val)
                    c2.border = BORDER
                    if std_val is not None:
                        c2.number_format = '0.0000'
                    col_idx += 2
                r += 1
            r += 1  # 空行分隔
    return r


def write_c1_c2_iter_block(ws, start_row, title, pair_results):
    """写入 C1/C2 拟合结果（7 组 Delta）。"""
    ws.cell(row=start_row, column=1, value=title).font = TITLE_FONT
    r = start_row + 1
    headers = ['Delta对', '参数', 'MC', 'IC']
    for j, h in enumerate(headers):
        style_header(ws.cell(row=r, column=j + 1, value=h))
    r += 1
    for pair_name, results in pair_results:
        for param_name, key in [('C1', 'C1'), ('C2', 'C2'), ('RMSE', 'RMSE'), ('细胞数', 'n_cells')]:
            ws.cell(row=r, column=1, value=pair_name).border = BORDER
            ws.cell(row=r, column=2, value=param_name).border = BORDER
            for ci, label in enumerate(['MC', 'IC']):
                val = results.get(label, {}).get(key, 0)
                if key == 'n_cells':
                    c = ws.cell(row=r, column=3 + ci, value=int(val))
                else:
                    c = ws.cell(row=r, column=3 + ci, value=val)
                    c.number_format = '0.000000'
                c.border = BORDER
            r += 1
    return r + 1


# ---------- 主处理 ----------
def process_iteration(base, constraint, all_c1_c2_results):
    out_path = os.path.join(base, f'{constraint}_iteration_汇总.xlsx')
    wb = Workbook()
    wb.remove(wb.active)

    runs = list(range(1, 11))
    ell = read_iter_ellipse(base, constraint, runs)
    ea = read_iter_edgeangle(base, constraint, runs)
    mm = read_iter_mema(base, constraint, runs)

    # MC/IC 分离（每个迭代点）
    ell_split = {}
    for it in ITERATIONS:
        mc, ic = split_mc_ic(ell[it])
        ell_split[it] = {'MC': mc, 'IC': ic}

    # 三分类（iteration 0）
    ell0_mc, ell0_ncmc, ell0_ic3 = split_three_way(ell[0])

    titles_mc_ic = []
    for it in ITERATIONS:
        titles_mc_ic += [f'iter{it}_MC', f'iter{it}_IC']

    # ---- ellipse 原始 ----
    ws = wb.create_sheet('ellipse')
    iter_hdr = ['run', 'No.', 'a', 'b', 'mn', 'n', 'P', 'A', '当前细胞层']
    groups = []
    for it in ITERATIONS:
        groups += [ell_split[it]['MC'], ell_split[it]['IC']]
    headers_list = [iter_hdr] * len(groups)
    # 列名加迭代后缀
    for i, it in enumerate(ITERATIONS):
        suffix = f'_iter{it}'
        headers_list[i * 2] = [c + suffix for c in iter_hdr]
        headers_list[i * 2 + 1] = [c + suffix for c in iter_hdr]
    write_raw_sheet(ws, groups, titles_mc_ic, headers_list, iter_hdr)

    # ---- 计算 7 组 Delta ----
    all_deltas = []
    pair_c1_c2 = []
    for pair_name, to_it, from_it in DELTA_PAIRS:
        delta_df = compute_delta_iter(ell[from_it], ell[to_it], from_it, to_it)
        all_deltas.append((pair_name, delta_df))
        # C1/C2 拟合
        results = fit_c1_c2_iter(ell[from_it], ell[to_it], from_it, to_it)
        pair_c1_c2.append((pair_name, results))

    # ---- delta_detail 工作表 ----
    ws = wb.create_sheet('delta_detail')
    write_delta_detail_sheet(ws, all_deltas, ell[0]['当前细胞层'].values)

    # ---- ellipse_统计 ----
    ws = wb.create_sheet('ellipse_统计')
    r = 1
    # Block A-1: mn 三分类（iteration 0）
    dd_mn3 = {'MC': ell0_mc, 'NC-MC': ell0_ncmc, 'IC≥3': ell0_ic3}
    r = write_overall(ws, r, 'Block A-1: mn 统计 (iter0, MC/NC-MC/IC≥3)', ['mn'], dd_mn3)
    # Block A-2: n, P, A 二分类（5 个迭代点）
    dd_all = {}
    for it in ITERATIONS:
        dd_all[f'iter{it}_MC'] = ell_split[it]['MC']
        dd_all[f'iter{it}_IC'] = ell_split[it]['IC']
    r = write_overall(ws, r, 'Block A-2: n, P, A 统计 (5迭代, MC/IC)', ['n', 'P', 'A'], dd_all)
    # Block B-0: n边形数量和频率（iter0, MC/IC）
    dd_init = {'MC': ell_split[0]['MC'], 'IC': ell_split[0]['IC']}
    r = write_n_dist_separate(ws, r, 'Block B-0: n边形数量和频率 (iter0)', dd_init)
    # Block B-1: mn 子区（iter0, 三分类）
    r = write_by_n(ws, r, 'Block B-1: mn 子区 (iter0, 三分类)', dd_mn3, 'mn')
    # Block B-2/3: P/A 子区（5 迭代点）
    r = write_by_n(ws, r, 'Block B-2: P 子区 (5迭代, 按n)', dd_all, 'P')
    r = write_by_n(ws, r, 'Block B-3: A 子区 (5迭代, 按n)', dd_all, 'A')
    # Block C: Delta 按 n 分组统计 (MC/IC → delta值 → delta对)
    r = write_delta_block_c(ws, r, 'Block C: Delta 统计 (按n, MC/IC × delta值 × delta对)', all_deltas)
    # Block D: C1/C2 拟合
    r = write_c1_c2_iter_block(ws, r, 'Block D: C1/C2 拟合 (7组Delta)', pair_c1_c2)

    # 存储 C1/C2 结果
    all_c1_c2_results[constraint] = pair_c1_c2

    # ---- edgeAngle 原始 ----
    ws = wb.create_sheet('edgeAngle')
    ea_split = {}
    for it in ITERATIONS:
        mc, ic = split_mc_ic(ea[it])
        ea_split[it] = {'MC': mc, 'IC': ic}
    ea_hdr = ['run', 'No.', 'n', 'mn', '内角', '夹边1', '当前细胞层']
    groups_ea = []
    for it in ITERATIONS:
        groups_ea += [ea_split[it]['MC'], ea_split[it]['IC']]
    headers_ea = [ea_hdr] * len(groups_ea)
    for i, it in enumerate(ITERATIONS):
        suffix = f'_iter{it}'
        headers_ea[i * 2] = [c + suffix for c in ea_hdr]
        headers_ea[i * 2 + 1] = [c + suffix for c in ea_hdr]
    write_raw_sheet(ws, groups_ea, titles_mc_ic, headers_ea, ea_hdr)

    # ---- edgeAngle_统计 ----
    ws = wb.create_sheet('edgeAngle_统计')
    dd_ea = {}
    for it in ITERATIONS:
        dd_ea[f'iter{it}_MC'] = ea_split[it]['MC']
        dd_ea[f'iter{it}_IC'] = ea_split[it]['IC']
    dd_ea_deg = {}
    for k, v in dd_ea.items():
        v2 = v.copy()
        v2['内角'] = v2['内角'] * 180.0 / np.pi
        dd_ea_deg[k] = v2
    r = 1
    r = write_overall(ws, r, 'Block A: 整体统计 (内角[度], 夹边1)', ['内角', '夹边1'], dd_ea_deg)
    r = write_dist_separate(ws, r, 'Block B: 边长分布 (夹边1)', '夹边1', dd_ea,
                            length_bin, LENGTH_LABELS)
    r = write_dist_separate(ws, r, 'Block C: 角度分布 (内角)', '内角', dd_ea,
                            angle_bin, ANGLE_LABELS, is_angle_from_rad=True)

    # ---- ME_MA 原始 ----
    ws = wb.create_sheet('ME_MA')
    mm_split = {}
    for it in ITERATIONS:
        mc = mm[it][mm[it]['当前细胞层'] == 1].reset_index(drop=True)
        mm_split[it] = mc
    mm_hdr = ['run', 'No.', 'ME', 'MA1', 'MA2', '当前细胞层']
    groups_mm = [mm_split[it] for it in ITERATIONS]
    titles_mm = [f'iter{it}_MC' for it in ITERATIONS]
    headers_mm = [mm_hdr] * len(groups_mm)
    for i, it in enumerate(ITERATIONS):
        headers_mm[i] = [c + f'_iter{it}' for c in mm_hdr]
    write_raw_sheet(ws, groups_mm, titles_mm, headers_mm, mm_hdr)

    # ---- ME_MA_统计 ----
    ws = wb.create_sheet('ME_MA_统计')
    dd_mm = {}
    for it in ITERATIONS:
        dd_mm[f'iter{it}_MC'] = mm_split[it]
    r = 1
    r = write_overall(ws, r, 'Block A: 整体统计 (ME, MA1, MA2)', ['ME', 'MA1', 'MA2'], dd_mm)
    r = write_dist_separate(ws, r, 'Block B: 边长分布 (ME)', 'ME', dd_mm,
                            length_bin, LENGTH_LABELS)
    r = write_dist_separate(ws, r, 'Block C: 角度分布 (MA1)', 'MA1', dd_mm,
                            angle_bin, ANGLE_LABELS)
    r = write_dist_separate(ws, r, 'Block D: 角度分布 (MA2)', 'MA2', dd_mm,
                            angle_bin, ANGLE_LABELS)

    wb.save(out_path)
    print(f'Saved: {out_path}')


def generate_iteration_c1_summary(base, all_c1_c2_results):
    """生成跨约束 C1/C2 对比文件（合并为一个 sheet，用标题分区）。"""
    out_path = os.path.join(base, 'iteration_C1_C2_汇总对比.xlsx')
    wb = Workbook()
    wb.remove(wb.active)

    # 按约束类型分组
    groups = {}
    for constraint, pair_results in all_c1_c2_results.items():
        group_name = '有120约束' if 'with120' in constraint else '无120约束'
        groups.setdefault(group_name, {})[constraint] = pair_results

    # Sheet 1: C1/C2 汇总（所有约束放一起，用标题分区）
    ws = wb.create_sheet('C1_C2汇总')
    headers = ['Delta对', 'C1_MC', 'C1_IC', 'C2_MC', 'C2_IC',
               'RMSE_MC', 'RMSE_IC', '细胞数_MC', '细胞数_IC']
    for cc in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(cc)].width = 14

    r = 1
    for group_name, constraint_dict in groups.items():
        # 分区标题
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
        c = ws.cell(row=r, column=1, value=f'===== {group_name} =====')
        c.font = TITLE_FONT
        c.fill = GROUP_FILL
        c.alignment = CENTER
        r += 1
        # 表头
        for j, h in enumerate(headers):
            style_header(ws.cell(row=r, column=j + 1, value=h))
        r += 1
        for constraint, pair_results in constraint_dict.items():
            for pair_name, results in pair_results:
                ws.cell(row=r, column=1, value=pair_name).border = BORDER
                for ci, key in enumerate(['C1', 'C2', 'RMSE']):
                    for label_idx, label in enumerate(['MC', 'IC']):
                        val = results.get(label, {}).get(key, 0)
                        c = ws.cell(row=r, column=2 + ci * 2 + label_idx, value=val)
                        c.border = BORDER
                        c.number_format = '0.000000'
                for label_idx, label in enumerate(['MC', 'IC']):
                    val = results.get(label, {}).get('n_cells', 0)
                    c = ws.cell(row=r, column=8 + label_idx, value=int(val))
                    c.border = BORDER
                r += 1
            r += 1

    # Sheet 2: 拟合明细
    ws2 = wb.create_sheet('拟合明细')
    r = 1
    for group_name, constraint_dict in groups.items():
        ws2.cell(row=r, column=1, value=f'===== {group_name} =====').font = TITLE_FONT
        r += 1
        for constraint, pair_results in constraint_dict.items():
            for pair_name, results in pair_results:
                ws2.cell(row=r, column=1, value=f'{constraint} - Δ({pair_name})').font = TITLE_FONT
                r += 1
                for label in ['MC', 'IC']:
                    det = results.get(label, {}).get('details', pd.DataFrame())
                    if len(det) == 0:
                        continue
                    ws2.cell(row=r, column=1, value=f'类型: {label}').font = Font(bold=True)
                    r += 1
                    cols = ['run', 'No.', 'n', 'ΔA', 'A_from', 'An', '(1-A/An)', 'ΔA_pred', '残差']
                    for j, h in enumerate(cols):
                        style_header(ws2.cell(row=r, column=j + 1, value=h))
                    r += 1
                    for _, row_data in det.iterrows():
                        vals = [row_data.get(c, '') for c in cols]
                        for j, val in enumerate(vals):
                            cell = ws2.cell(row=r, column=j + 1, value=val)
                            cell.border = BORDER
                            if isinstance(val, float):
                                cell.number_format = '0.000000'
                        r += 1
                    r += 1
                r += 1
        r += 1

    wb.save(out_path)
    print(f'Saved: {out_path}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(
        description='Iteration 数据汇总：分析退火迭代的细胞数据变化。'
    )
    parser.add_argument('base_dir', help='iteration 根目录路径')
    parser.add_argument('--constraints', nargs='*', default=None,
                        help='可选：指定约束文件夹名')
    args = parser.parse_args()

    base_dir = args.base_dir
    if not os.path.isdir(base_dir):
        print(f"Error: 目录不存在: {base_dir}")
        sys.exit(1)

    # 发现约束文件夹
    if args.constraints:
        constraints = args.constraints
    else:
        constraints = [d for d in os.listdir(base_dir)
                       if os.path.isdir(os.path.join(base_dir, d)) and 'run1' in os.listdir(os.path.join(base_dir, d))]

    print(f"扫描目录: {base_dir}")
    print(f"发现 {len(constraints)} 个约束文件夹: {constraints}")
    print("-" * 50)

    all_c1_c2_results = {}
    for constraint in sorted(constraints):
        print(f"Processing: {constraint} ...")
        process_iteration(base_dir, constraint, all_c1_c2_results)
        print(f"  Done: {constraint}")

    print("-" * 50)
    if all_c1_c2_results:
        generate_iteration_c1_summary(base_dir, all_c1_c2_results)
    print("-" * 50)
    print("ALL DONE！")
