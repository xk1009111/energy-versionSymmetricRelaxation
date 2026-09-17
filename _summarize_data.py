# -*- coding: utf-8 -*-
"""
_summarize_data.py — 批量退火数据汇总（纯 openpyxl + numpy）。

通用配置化版本：所有实验参数、列名、分类规则、分布 bin 都集中在 CONFIG 里。
改实验只需改 CONFIG，不用动主逻辑。

数据结构: {base}/{net}/{lam_dir}/run{r}/{iter}_{phase}_{type}.xlsx
输出: 每个 net 一个汇总 .xlsx + 一个跨 net 的 C1_C2 汇总对比 .xlsx

运行: python _summarize_data.py
"""
import os, sys
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================================
# CONFIG — 改实验参数只改这里
# ============================================================================
CONFIG = {
    # ---- 路径 ----
    'base_dir': 'data_fixed',           # 数据根目录
    'runs': list(range(1, 11)),          # run 编号列表

    # ---- λ 配置 ----
    'lam_dirs': ['L0', 'L1', 'L10', 'L100'],  # λ 子目录名
    'lam_values': [0, 1, 10, 100],            # λ 数值（与 lam_dirs 一一对应）

    # ---- 迭代阶段 ----
    'iterations': [0, 1],                # [init, final]
    'phase_map': {0: 'initial', 1: 'final'},

    # ---- 文件类型（annealing GUI 导出的三种 xlsx）----
    'file_types': ['ellipse', 'edgeAngle', 'ME_MA'],

    # ---- 分类规则 ----
    'mc_layer': 1,                       # 边缘细胞层号
    'three_way': ['MC', 'NC-MC', 'IC≥3'],  # mn 三分类名（层1 / 层2 / 层≥3）

    # ---- 分布 bin ----
    'length_bins': [0.1, 0.3, 0.5, 0.7, 0.9, 1.1, 1.3, 1.5, 1.7, 1.9, float('inf')],
    'length_labels': ['<0.1', '0.1-0.3', '0.3-0.5', '0.5-0.7', '0.7-0.9',
                      '0.9-1.1', '1.1-1.3', '1.3-1.5', '1.5-1.7', '1.7-1.9', '>1.9'],
    'angle_bins': [30, 50, 70, 90, 110, 130, 150, 170, 180],
    'angle_labels': ['<30°', '30-50°', '50-70°', '70-90°', '90-110°',
                     '110-130°', '130-150°', '150-170°', '170-180°'],

    # ---- 列名映射（xlsx 中文列名 -> 内部通用 key）----
    'col_map': {
        # ellipse
        'run': 'run',
        '细胞序号': 'No.',
        '长半轴': 'a',
        '短半轴': 'b',
        '相邻细胞边数和': 'mn',
        '细胞边数': 'n',
        '细胞周长': 'P',
        '细胞面积': 'A',
        '当前细胞层': 'layer',
        # edgeAngle
        '边数': 'n_ea',
        '内角': 'alpha',
        '夹边1': 'edge1',
        # ME_MA
        '边缘边长(ME)': 'ME',
        '边缘角1(MA1)': 'MA1',
        '边缘角2(MA2)': 'MA2',
    },
}

# ============================================================================
# 样式常量
# ============================================================================
THIN = Side(style='thin', color='000000')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
HEADER_FONT = Font(bold=True, size=11)
TITLE_FONT = Font(bold=True, size=12)
HEADER_FILL = PatternFill(start_color='B8CCE4', end_color='B8CCE4', fill_type='solid')
GROUP_FILL = PatternFill(start_color='D5E8D4', end_color='D5E8D4', fill_type='solid')
DELTA_FILL = PatternFill(start_color='F8CECC', end_color='F8CECC', fill_type='solid')

def style_header(cell):
    cell.font = HEADER_FONT; cell.fill = HEADER_FILL
    cell.alignment = CENTER; cell.border = BORDER

def _write_cell(ws, r, c, val, fmt=None):
    cell = ws.cell(row=r, column=c, value=val)
    cell.border = BORDER
    if fmt and isinstance(val, (int, float)): cell.number_format = fmt
    return cell

# ============================================================================
# 工具函数
# ============================================================================
def mean_std(vals):
    vals = [float(v) for v in vals if v is not None and not (isinstance(v, float) and np.isnan(v))]
    if len(vals) == 0: return None, None
    return float(np.mean(vals)), float(np.std(vals, ddof=1)) if len(vals) > 1 else 0.0

def group_by_n(rows, n_col='n'):
    d = {}
    for r in rows:
        n = int(r.get(n_col, 0))
        d.setdefault(n, []).append(r)
    return d

def length_bin(v, cfg):
    for i, b in enumerate(cfg['length_bins']):
        if v <= b: return i
    return len(cfg['length_bins']) - 1

def angle_bin(v, cfg):
    for i, b in enumerate(cfg['angle_bins']):
        if v <= b: return i
    return len(cfg['angle_bins']) - 1

# ============================================================================
# 分类
# ============================================================================
def mc_ic(rows, cfg):
    """二分类: MC=层1, IC=层≥2"""
    mc = [r for r in rows if int(r.get('layer', 0)) == cfg['mc_layer']]
    ic = [r for r in rows if int(r.get('layer', 0)) >= cfg['mc_layer'] + 1]
    return mc, ic

def three_way(rows, cfg):
    """三分类: MC=层1, NC-MC=层2, IC≥3"""
    mc = [r for r in rows if int(r.get('layer', 0)) == cfg['mc_layer']]
    ncmc = [r for r in rows if int(r.get('layer', 0)) == cfg['mc_layer'] + 1]
    ic3 = [r for r in rows if int(r.get('layer', 0)) >= cfg['mc_layer'] + 2]
    return mc, ncmc, ic3

# ============================================================================
# 发现 & 读取
# ============================================================================
def discover_nets(cfg):
    base = cfg['base_dir']
    return sorted([d for d in os.listdir(base)
                   if os.path.isdir(os.path.join(base, d))
                   and any(os.path.isdir(os.path.join(base, d, l)) for l in cfg['lam_dirs'])])

def read_xlsx_rows(path, cfg):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows: return []
    header = list(rows[0])
    cm = cfg['col_map']
    # 把中文列名映射为内部 key
    return [{cm.get(h, h): row[j] for j, h in enumerate(header)} for row in rows[1:]]

def collect(base, net, lam, runs, it, phase, typ, cfg):
    frames = []
    for run in runs:
        path = os.path.join(base, net, lam, f'run{run}',
                            f'{it}_{phase}_{typ}.xlsx')
        if not os.path.exists(path): continue
        for row in read_xlsx_rows(path, cfg):
            row['run'] = run
            frames.append(row)
    return frames

# ============================================================================
# Delta / C1C2
# ============================================================================
def compute_delta_rows(from_rows, to_rows):
    from_map = {(int(r['run']), int(r['No.'])): r for r in from_rows}
    to_map = {(int(r['run']), int(r['No.'])): r for r in to_rows}
    keys = set(from_map.keys()) & set(to_map.keys())
    out = []
    for key in keys:
        fr, tr = from_map[key], to_map[key]
        Pf = float(fr['P']); Pt = float(tr['P'])
        Af = float(fr['A']); At = float(tr['A'])
        nf = int(fr.get('n', 0))
        layer = int(fr.get('layer', 0))
        SI_f = Pf / np.sqrt(max(Af, 1e-12))
        SI_t = Pt / np.sqrt(max(At, 1e-12))
        out.append({
            'run': key[0], 'No.': key[1], 'n': nf,
            'ΔP': Pt - Pf, 'ΔA': At - Af,
            'ΔSI': SI_t - SI_f, 'ΔEavg': (Pt - Pf) / max(nf, 1),
            'layer': layer, 'A_from': Af,
        })
    return out

def fit_c1_c2(delta_rows, cfg):
    mc_layer = cfg['mc_layer']
    results = {}
    for (label, n_offset, mask_fn) in [
        ('MC', 5, lambda r: int(r.get('layer', 0)) == mc_layer),
        ('IC', 6, lambda r: int(r.get('layer', 0)) >= mc_layer + 1)
    ]:
        sub = [r for r in delta_rows if mask_fn(r)]
        if not sub:
            results[label] = {'C1': 0, 'C2': 0, 'RMSE': 0, 'n_cells': 0, 'An': {}, 'details': []}
            continue
        An = {}
        for r in sub:
            nv = int(r['n'])
            An.setdefault(nv, []).append(float(r['A_from']))
        An = {k: float(np.mean(v)) for k, v in An.items()}
        n_arr = np.array([int(r['n']) for r in sub], dtype=float)
        da_arr = np.array([float(r['ΔA']) for r in sub], dtype=float)
        a_arr = np.array([float(r['A_from']) for r in sub], dtype=float)
        dn = n_arr - n_offset
        mask = dn != 0
        c1 = float(np.sum(dn[mask] * da_arr[mask]) / np.sum(dn[mask] ** 2)) if mask.any() else 0.0
        correction = np.array([1 - a_arr[i] / An.get(int(n_arr[i]), 1.0) for i in range(len(sub))])
        best_c2, best_rmse = 0.0, float('inf')
        for c2_val in np.arange(0.0, 3.001, 0.01):
            pred = c1 * dn + c2_val * correction
            rmse = float(np.sqrt(np.mean((da_arr - pred) ** 2)))
            if rmse < best_rmse:
                best_rmse, best_c2 = rmse, float(c2_val)
        results[label] = {'C1': c1, 'C2': best_c2, 'RMSE': best_rmse,
                           'n_cells': len(sub), 'An': An, 'details': sub}
    return results

# ============================================================================
# 写入辅助
# ============================================================================
def _write_dist_alt(ws, start_r, title, groups_dict, col_key, bin_fn, labels, cfg, is_angle=False):
    ws.cell(row=start_r, column=1, value=title).font = TITLE_FONT
    r = start_r + 1
    n_bins = len(labels)
    for gname, rows in groups_dict.items():
        ws.cell(row=r, column=1, value=gname).font = Font(bold=True)
        r += 1
        headers = ['统计量'] + labels + ['总数'] + ['均值'] + ['标准差']
        for j, h in enumerate(headers): style_header(ws.cell(row=r, column=j+1, value=h))
        r += 1
        vals_raw = [row.get(col_key) for row in rows if row.get(col_key) is not None]
        if is_angle:
            vals = [float(v) * 180.0 / np.pi if isinstance(v, float) and v < 6.5 else float(v)
                    for v in vals_raw]
        else:
            vals = [float(v) for v in vals_raw]
        bins = [0] * n_bins
        for v in vals:
            idx = bin_fn(v, cfg)
            if 0 <= idx < n_bins: bins[idx] += 1
        total = len(vals)
        for row_label, data in [('数量', bins)]:
            _write_cell(ws, r, 1, row_label)
            for j, v in enumerate(data): _write_cell(ws, r, 2 + j, v)
            _write_cell(ws, r, 2 + n_bins, total)
            m, s = mean_std(vals)
            _write_cell(ws, r, 2 + n_bins + 1, m, '0.0000')
            _write_cell(ws, r, 2 + n_bins + 2, s, '0.0000')
            r += 1
        freq = [f'{b} ({b/total:.1%})' if total else '0' for b in bins]
        _write_cell(ws, r, 1, '频率')
        for j, v in enumerate(freq): _write_cell(ws, r, 2 + j, v)
        _write_cell(ws, r, 2 + n_bins, total)
        r += 1
        r += 1
    return r + 1

def _write_delta_block_c(ws, start_r, title, all_pairs_data):
    ws.cell(row=start_r, column=1, value=title).font = TITLE_FONT
    r = start_r + 1
    delta_metrics = ['ΔP', 'ΔA', 'ΔSI', 'ΔEavg']
    pair_names = [p[0] for p in all_pairs_data]
    for (mc_label, rows_list) in [('MC', [p[1] for p in all_pairs_data]),
                                   ('IC', [p[2] for p in all_pairs_data])]:
        for metric in delta_metrics:
            ws.cell(row=r, column=1, value=f'{mc_label} - {metric}').font = Font(bold=True)
            r += 1
            all_n = set()
            for rows in rows_list:
                for row in rows:
                    n = row.get('n')
                    if n is not None: all_n.add(int(n))
            all_n = sorted(all_n)
            headers = ['n值']
            for pn in pair_names: headers += [f'{pn}_均值', f'{pn}_标准差']
            for j, h in enumerate(headers): style_header(ws.cell(row=r, column=j+1, value=h))
            r += 1
            for nv in all_n:
                _write_cell(ws, r, 1, nv)
                col = 2
                for pi, rows in enumerate(rows_list):
                    sub = [row for row in rows if int(row.get('n', -1)) == nv]
                    m, s = mean_std([row.get(metric) for row in sub])
                    _write_cell(ws, r, col, m, '0.0000')
                    _write_cell(ws, r, col + 1, s, '0.0000')
                    col += 2
                r += 1
            r += 1
    return r + 1

def write_raw_block(ws, titles, data_lists, headers):
    n = len(titles); cols_per = len(headers)
    for i in range(n):
        start_col = 1 + i * cols_per
        ws.merge_cells(start_row=1, start_column=start_col,
                       end_row=1, end_column=start_col + cols_per - 1)
        c = ws.cell(row=1, column=start_col, value=titles[i])
        c.font = HEADER_FONT; c.fill = DELTA_FILL; c.alignment = CENTER
        for cc in range(start_col, start_col + cols_per):
            ws.cell(row=1, column=cc).fill = DELTA_FILL
        for j, h in enumerate(headers):
            style_header(ws.cell(row=2, column=start_col + j, value=h))
        rows = data_lists[i]
        for ri, row_data in enumerate(rows):
            for j, h in enumerate(headers):
                val = row_data.get(h)
                c = ws.cell(row=3 + ri, column=start_col + j, value=val)
                c.border = BORDER
                if isinstance(val, float): c.number_format = '0.000000'
    ws.freeze_panes = 'A3'

def write_delta_detail_block(ws, all_pairs_data):
    delta_metrics = ['ΔP', 'ΔA', 'ΔSI', 'ΔEavg']
    pair_names = [p[0] for p in all_pairs_data]
    n_pairs = len(pair_names); cols_per_group = n_pairs * len(delta_metrics)
    col = 1
    for label, idx in [('MC', 1), ('IC', 2)]:
        merged = {}
        for pi, (pair_name, mc_r, ic_r) in enumerate(all_pairs_data):
            target = mc_r if label == 'MC' else ic_r
            for r in target:
                key = (r.get('run'), r.get('No.'))
                merged.setdefault(key, {}).update(r)
                for dm in delta_metrics:
                    merged[key].update({f'{dm}_{pair_name}': r.get(dm)})
        sorted_rows = [merged[k] for k in sorted(merged.keys())]
        start_col = col
        ws.merge_cells(start_row=1, start_column=start_col,
                       end_row=1, end_column=start_col + cols_per_group - 1)
        c = ws.cell(row=1, column=start_col, value=label)
        c.font = HEADER_FONT; c.fill = DELTA_FILL; c.alignment = CENTER
        for cc in range(start_col, start_col + cols_per_group):
            ws.cell(row=1, column=cc).fill = DELTA_FILL
        for mi, metric in enumerate(delta_metrics):
            metric_start = start_col + mi * n_pairs
            ws.merge_cells(start_row=2, start_column=metric_start,
                           end_row=2, end_column=metric_start + n_pairs - 1)
            c2 = ws.cell(row=2, column=metric_start, value=metric)
            c2.font = HEADER_FONT; c2.fill = DELTA_FILL; c2.alignment = CENTER
            for cc in range(metric_start, metric_start + n_pairs):
                ws.cell(row=2, column=cc).fill = DELTA_FILL
            for j, pn in enumerate(pair_names):
                style_header(ws.cell(row=3, column=metric_start + j, value=pn))
                ws.cell(row=3, column=metric_start + j).fill = DELTA_FILL
            col = metric_start + n_pairs
        for ri, row_data in enumerate(sorted_rows):
            for mi, metric in enumerate(delta_metrics):
                for j, pn in enumerate(pair_names):
                    key = f'{metric}_{pn}'
                    val = row_data.get(key)
                    c = ws.cell(row=4 + ri, column=start_col + mi * n_pairs + j,
                                value=float(val) if val is not None else None)
                    c.border = BORDER
                    if val is not None: c.number_format = '0.000000'
    ws.freeze_panes = 'A4'

# ============================================================================
# 主流程
# ============================================================================
def process_net(base, net, cfg, all_c1):
    out_path = os.path.join(base, f'{net}_iteration_汇总.xlsx')
    runs = cfg['runs']
    lam_dirs = cfg['lam_dirs']
    its = cfg['iterations']
    phase_map = cfg['phase_map']

    # ---- 加载 ----
    ell, ea, mm = {}, {}, {}
    for lam in lam_dirs:
        for it in its:
            phase = phase_map[it]
            ell.setdefault(lam, {})[it] = collect(base, net, lam, runs, it, phase, 'ellipse', cfg)
            ea.setdefault(lam, {})[it] = collect(base, net, lam, runs, it, phase, 'edgeAngle', cfg)
            mm.setdefault(lam, {})[it] = collect(base, net, lam, runs, it, phase, 'ME_MA', cfg)

    # ---- 分类 ----
    mc_ell, ic_ell, tway = {}, {}, {}
    mc_ea, ic_ea = {}, {}
    mc_mm = {}
    for lam in lam_dirs:
        mc_ell[lam], ic_ell[lam] = {}, {}
        mc_ea[lam], ic_ea[lam] = {}, {}
        mc_mm[lam] = {}
        for it in its:
            mc_ell[lam][it], ic_ell[lam][it] = mc_ic(ell[lam][it], cfg)
            mc_ea[lam][it], ic_ea[lam][it] = mc_ic(ea[lam][it], cfg)
            mm_mc = [r for r in mm[lam][it] if int(r.get('layer', 0)) == cfg['mc_layer']]
            mc_mm[lam][it] = mm_mc
        tway[lam] = three_way(ell[lam][0], cfg)

    # init 合并池
    ell_init_pool = []
    ea_init_pool = []
    mm_init_pool = []
    for lam in lam_dirs:
        ell_init_pool += ell[lam][0]
        ea_init_pool += ea[lam][0]
        mm_init_pool += mc_mm[lam][0]
    tway_pool = three_way(ell_init_pool, cfg)

    # ---- Delta + C1C2 ----
    all_pairs_data = []
    c1_c2_results = {}
    for lam in lam_dirs:
        delta = compute_delta_rows(ell[lam][0], ell[lam][1])
        mc_d, ic_d = mc_ic(delta, cfg)
        all_pairs_data.append((lam, mc_d, ic_d))
        c1_c2_results[lam] = fit_c1_c2(delta, cfg)

    # ---- 写 workbook ----
    wb = Workbook(); wb.remove(wb.active)

    # ===== ellipse 原始 =====
    ws = wb.create_sheet('ellipse')
    titles_raw, data_raw = [], []
    for lam in lam_dirs:
        for it in its:
            titles_raw += [f'{lam}_{phase_map[it]}_MC', f'{lam}_{phase_map[it]}_IC']
            data_raw += [mc_ell[lam][it], ic_ell[lam][it]]
    write_raw_block(ws, titles_raw, data_raw,
                    ['run', 'No.', 'a', 'b', 'mn', 'n', 'P', 'A', 'layer'])

    # ===== delta_detail =====
    ws = wb.create_sheet('delta_detail')
    write_delta_detail_block(ws, all_pairs_data)

    # ===== ellipse_统计 =====
    ws = wb.create_sheet('ellipse_统计')
    r = 1
    LBL = cfg['three_way']

    # Block A-1: mn 三分类（init 合并）
    ws.cell(row=r, column=1, value=f'Block A-1: mn 三分类 (MC/NC-MC/IC≥3, init 合并{len(lam_dirs)}λ×runs)').font = TITLE_FONT
    r += 1
    for j, h in enumerate(['分类', '均值', '标准差']): style_header(ws.cell(row=r, column=j + 1, value=h))
    r += 1
    for cat_idx, cat_name in enumerate(LBL):
        m, s = mean_std([row.get('mn') for row in tway_pool[cat_idx]])
        _write_cell(ws, r, 1, cat_name)
        _write_cell(ws, r, 2, m, '0.0000')
        _write_cell(ws, r, 3, s, '0.0000')
        r += 1
    r += 1

    # Block A-2: n/P/A（init 合并, final × λ）
    for col_key, short in [('n', 'n'), ('P', 'P'), ('A', 'A')]:
        ws.cell(row=r, column=1, value=f'Block A-2-{short}: {short}').font = TITLE_FONT
        r += 1
        headers_a2 = ['分类', 'init_均值', 'init_标准差']
        for lam in lam_dirs:
            headers_a2 += [f'{lam}_final_均值', f'{lam}_final_标准差']
        for j, h in enumerate(headers_a2): style_header(ws.cell(row=r, column=j + 1, value=h))
        r += 1
        for sub_label, mask_fn in [('MC', lambda rr: int(rr.get('layer', 0)) == cfg['mc_layer']),
                                    ('IC', lambda rr: int(rr.get('layer', 0)) >= cfg['mc_layer'] + 1)]:
            _write_cell(ws, r, 1, sub_label)
            col = 2
            m, s = mean_std([row.get(col_key) for row in ell_init_pool if mask_fn(row)])
            _write_cell(ws, r, col, m, '0.0000'); _write_cell(ws, r, col + 1, s, '0.0000')
            col += 2
            for lam in lam_dirs:
                rows = [row for row in ell[lam][1] if mask_fn(row)]
                m, s = mean_std([row.get(col_key) for row in rows])
                _write_cell(ws, r, col, m, '0.0000'); _write_cell(ws, r, col + 1, s, '0.0000')
                col += 2
            r += 1
        r += 1

    # Block B-0: n 边形分布（init 合并）
    ws.cell(row=r, column=1, value='Block B-0: n 边形分布 (init 合并, MC/IC)').font = TITLE_FONT
    r += 1
    for sub_label, mask_fn in [('MC', lambda rr: int(rr.get('layer', 0)) == cfg['mc_layer']),
                                ('IC', lambda rr: int(rr.get('layer', 0)) >= cfg['mc_layer'] + 1)]:
        rows = [row for row in ell_init_pool if mask_fn(row)]
        ws.cell(row=r, column=1, value=sub_label).font = Font(bold=True)
        r += 1
        headers = ['n值', '数量', '频率']
        for j, h in enumerate(headers): style_header(ws.cell(row=r, column=j+1, value=h))
        r += 1
        by_n = group_by_n(rows)
        total = sum(len(v) for v in by_n.values())
        for nv in sorted(by_n.keys()):
            cnt = len(by_n[nv])
            _write_cell(ws, r, 1, nv)
            _write_cell(ws, r, 2, cnt)
            _write_cell(ws, r, 3, f'{cnt/total:.1%}' if total else '0')
            r += 1
        r += 1

    # Block B: P/A 按 n 分组（全 λ × 全阶段 × MC/IC）
    for col_key, short in [('P', 'P'), ('A', 'A')]:
        ws.cell(row=r, column=1, value=f'Block B-{short}: {short} 按 n 分组').font = TITLE_FONT
        r += 1
        for lam in lam_dirs:
            for it in its:
                for sub_label, rows in [('MC', mc_ell[lam][it]), ('IC', ic_ell[lam][it])]:
                    ws.cell(row=r, column=1, value=f'{lam}_{phase_map[it]}_{sub_label}').font = Font(bold=True)
                    r += 1
                    headers = ['n值', f'{short}_均值', f'{short}_标准差']
                    for j, h in enumerate(headers): style_header(ws.cell(row=r, column=j+1, value=h))
                    r += 1
                    by_n = group_by_n(rows)
                    for nv in sorted(by_n.keys()):
                        vals = [float(row[col_key]) for row in by_n[nv] if row.get(col_key) is not None]
                        m, s = mean_std(vals)
                        _write_cell(ws, r, 1, nv)
                        _write_cell(ws, r, 2, m, '0.0000')
                        _write_cell(ws, r, 3, s, '0.0000')
                        r += 1
                    r += 1
            r += 1

    # Block C: Delta 按 n
    r = _write_delta_block_c(ws, r, 'Block C: Delta 按 n', all_pairs_data)

    # Block D: C1/C2
    ws.cell(row=r, column=1, value='Block D: C1/C2 拟合 (final-initial)').font = TITLE_FONT
    r += 1
    headers = ['λ', '参数', 'MC', 'IC']
    for j, h in enumerate(headers): style_header(ws.cell(row=r, column=j+1, value=h))
    r += 1
    for lam in lam_dirs:
        results = c1_c2_results[lam]
        for param_name, key in [('C1', 'C1'), ('C2', 'C2'), ('RMSE', 'RMSE'), ('细胞数', 'n_cells')]:
            _write_cell(ws, r, 1, lam)
            _write_cell(ws, r, 2, param_name)
            for ci, label in enumerate(['MC', 'IC']):
                val = results.get(label, {}).get(key, 0)
                fmt = '0.000000' if isinstance(val, float) else None
                _write_cell(ws, r, 3 + ci, int(val) if key == 'n_cells' else val, fmt)
            r += 1
    r += 1

    all_c1[net] = c1_c2_results

    # ===== edgeAngle 原始 =====
    ws = wb.create_sheet('edgeAngle')
    titles_ea, data_ea = [], []
    for lam in lam_dirs:
        for it in its:
            titles_ea += [f'{lam}_{phase_map[it]}_MC', f'{lam}_{phase_map[it]}_IC']
            data_ea += [mc_ea[lam][it], ic_ea[lam][it]]
    write_raw_block(ws, titles_ea, data_ea,
                    ['run', 'No.', 'n_ea', 'mn', 'alpha', 'edge1', 'layer'])

    # ===== edgeAngle_统计 =====
    ws = wb.create_sheet('edgeAngle_统计')
    r = 1
    for col_key, label in [('alpha', '内角[度]'), ('edge1', '夹边1')]:
        ws.cell(row=r, column=1, value=f'Block A: 整体 ({label}, init 合并)').font = TITLE_FONT
        r += 1
        headers = ['分类', 'init_均值', 'init_标准差']
        for lam in lam_dirs:
            headers += [f'{lam}_final_均值', f'{lam}_final_标准差']
        for j, h in enumerate(headers): style_header(ws.cell(row=r, column=j + 1, value=h))
        r += 1
        for sub_label, mask_fn in [('MC', lambda rr: int(rr.get('layer', 0)) == cfg['mc_layer']),
                                    ('IC', lambda rr: int(rr.get('layer', 0)) >= cfg['mc_layer'] + 1)]:
            _write_cell(ws, r, 1, sub_label)
            col = 2
            m, s = mean_std([row.get(col_key) for row in ea_init_pool if mask_fn(row)])
            _write_cell(ws, r, col, m, '0.0000'); _write_cell(ws, r, col + 1, s, '0.0000')
            col += 2
            for lam in lam_dirs:
                rows = [row for row in ea[lam][1] if mask_fn(row)]
                m, s = mean_std([row.get(col_key) for row in rows])
                _write_cell(ws, r, col, m, '0.0000'); _write_cell(ws, r, col + 1, s, '0.0000')
                col += 2
            r += 1
        r += 1
    # 分布块
    groups_ea = {'init_MC': [row for row in ea_init_pool if int(row.get('layer', 0)) == cfg['mc_layer']],
                 'init_IC': [row for row in ea_init_pool if int(row.get('layer', 0)) >= cfg['mc_layer'] + 1]}
    for lam in lam_dirs:
        for sub_label, rows in [('MC', mc_ea[lam][1]), ('IC', ic_ea[lam][1])]:
            groups_ea[f'{lam}_final_{sub_label}'] = rows
    r = _write_dist_alt(ws, r, 'Block B: 边长分布 (edge1)', groups_ea, 'edge1',
                        length_bin, cfg['length_labels'], cfg)
    r = _write_dist_alt(ws, r, 'Block C: 角度分布 (alpha, 已为度)', groups_ea, 'alpha',
                        angle_bin, cfg['angle_labels'], cfg)

    # ===== ME_MA 原始 =====
    ws = wb.create_sheet('ME_MA')
    titles_mm, data_mm = [], []
    for lam in lam_dirs:
        for it in its:
            titles_mm += [f'{lam}_{phase_map[it]}_MC']
            data_mm += [mc_mm[lam][it]]
    write_raw_block(ws, titles_mm, data_mm,
                    ['run', 'No.', 'ME', 'MA1', 'MA2', 'layer'])

    # ===== ME_MA_统计 =====
    ws = wb.create_sheet('ME_MA_统计')
    r = 1
    for col_key, label in [('ME', 'ME'), ('MA1', 'MA1'), ('MA2', 'MA2')]:
        ws.cell(row=r, column=1, value=f'Block A: 整体 ({label}, init 合并)').font = TITLE_FONT
        r += 1
        headers = ['init_均值', 'init_标准差']
        for lam in lam_dirs:
            headers += [f'{lam}_final_均值', f'{lam}_final_标准差']
        for j, h in enumerate(headers): style_header(ws.cell(row=r, column=j + 1, value=h))
        r += 1
        col = 1
        m, s = mean_std([row.get(col_key) for row in mm_init_pool])
        _write_cell(ws, r, col, m, '0.0000'); _write_cell(ws, r, col + 1, s, '0.0000')
        col += 2
        for lam in lam_dirs:
            rows = mc_mm[lam][1]
            m, s = mean_std([row.get(col_key) for row in rows])
            _write_cell(ws, r, col, m, '0.0000'); _write_cell(ws, r, col + 1, s, '0.0000')
            col += 2
        r += 1
    groups_mm = {'init_MC': mm_init_pool}
    for lam in lam_dirs:
        groups_mm[f'{lam}_final_MC'] = mc_mm[lam][1]
    r = _write_dist_alt(ws, r, 'Block B: ME 分布', groups_mm, 'ME',
                        length_bin, cfg['length_labels'], cfg)
    r = _write_dist_alt(ws, r, 'Block C: MA1 分布 (已为度)', groups_mm, 'MA1',
                        angle_bin, cfg['angle_labels'], cfg)
    r = _write_dist_alt(ws, r, 'Block D: MA2 分布 (已为度)', groups_mm, 'MA2',
                        angle_bin, cfg['angle_labels'], cfg)

    # 列宽
    for ws in wb.worksheets:
        for cc in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(cc)].width = 14

    wb.save(out_path)
    print(f'  Saved: {out_path}')

# ============================================================================
# 跨组对比
# ============================================================================
def generate_c1_summary(base, all_c1, cfg):
    out_path = os.path.join(base, 'iteration_C1_C2_汇总对比.xlsx')
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet('C1_C2汇总')
    headers = ['初始网络', 'λ', 'C1_MC', 'C1_IC', 'C2_MC', 'C2_IC',
               'RMSE_MC', 'RMSE_IC', '细胞数_MC', '细胞数_IC']
    for cc in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(cc)].width = 14
    r = 1
    for net in sorted(all_c1.keys()):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
        c = ws.cell(row=r, column=1, value=f'===== {net} =====')
        c.font = TITLE_FONT; c.fill = GROUP_FILL; c.alignment = CENTER
        r += 1
        for j, h in enumerate(headers): style_header(ws.cell(row=r, column=j + 1, value=h))
        r += 1
        for lam in cfg['lam_dirs']:
            results = all_c1[net][lam]
            _write_cell(ws, r, 1, net)
            _write_cell(ws, r, 2, lam)
            for ci, key in enumerate(['C1', 'C2', 'RMSE']):
                for li, label in enumerate(['MC', 'IC']):
                    val = results.get(label, {}).get(key, 0)
                    _write_cell(ws, r, 3 + ci * 2 + li, val, '0.000000')
            for li, label in enumerate(['MC', 'IC']):
                val = results.get(label, {}).get('n_cells', 0)
                _write_cell(ws, r, 9 + li, int(val))
            r += 1
        r += 1
    wb.save(out_path)
    print(f'Saved: {out_path}')

# ============================================================================
# 入口
# ============================================================================
if __name__ == '__main__':
    cfg = CONFIG
    nets = discover_nets(cfg)
    print(f'发现 {len(nets)} 个初始网络: {nets}')
    print(f'λ 配置: {cfg["lam_dirs"]}')
    print(f'Run 范围: {cfg["runs"][0]}-{cfg["runs"][-1]}')
    print('-' * 50)
    all_c1 = {}
    for net in sorted(nets):
        print(f'Processing: {net} ...')
        process_net(cfg['base_dir'], net, cfg, all_c1)
        print(f'  Done: {net}')
    print('-' * 50)
    if all_c1: generate_c1_summary(cfg['base_dir'], all_c1, cfg)
    print('-' * 50)
    print('ALL DONE!')
